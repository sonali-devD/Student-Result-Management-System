import os
import tkinter as tk
from tkinter import messagebox, ttk
import openpyxl

# Step 1: Excel File Setup
EXCEL_FILE = "student_records.xlsx"

# Create Excel file with headers if it does not exist
if not os.path.exists(EXCEL_FILE):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(["Roll No", "Name", "Class", "Python", "Java", "Math", "Dart", "SQL", "Total", "Percentage", "Result"])
    wb.save(EXCEL_FILE)


# Step 2: Helper function to clear the right output frame
def clear_right_frame():
    for widget in right_frame.winfo_children():
        widget.destroy()


# Step 3: Save Button Function
def save_data():
    s_name = name_entry.get().strip()
    s_roll = roll_entry.get().strip()
    s_class = class_entry.get().strip()
    p = sub1_entry.get().strip()
    j = sub2_entry.get().strip()
    m = sub3_entry.get().strip()
    d = sub4_entry.get().strip()
    s = sub5_entry.get().strip()

    # Validation: Check empty fields
    if not (s_name and s_roll and s_class and p and j and m and d and s):
        messagebox.showerror("Error", "Please fill all fields!")
        return

    # Convert inputs to numbers
    try:
        r_no = int(s_roll)
        m1, m2, m3, m4, m5 = float(p), float(j), float(m), float(d), float(s)
    except ValueError:
        messagebox.showerror("Error", "Roll No and Marks must be numbers!")
        return

    # Calculations
    total = m1 + m2 + m3 + m4 + m5
    percentage = total / 5.0
    
    # Pass/Fail Logic (35 passing mark per subject)
    if m1 >= 35 and m2 >= 35 and m3 >= 35 and m4 >= 35 and m5 >= 35:
        result = "Pass"
    else:
        result = "Fail"

    # Save to Excel File
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active

    # Check for duplicate roll number
    for row in ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]) == str(s_roll):
            messagebox.showerror("Error", "Roll Number already exists!")
            wb.close()
            return

    ws.append([r_no, s_name, s_class, m1, m2, m3, m4, m5, total, round(percentage, 2), result])
    wb.save(EXCEL_FILE)
    wb.close()

    # Clear input text boxes
    name_entry.delete(0, tk.END)
    roll_entry.delete(0, tk.END)
    class_entry.delete(0, tk.END)
    sub1_entry.delete(0, tk.END)
    sub2_entry.delete(0, tk.END)
    sub3_entry.delete(0, tk.END)
    sub4_entry.delete(0, tk.END)
    sub5_entry.delete(0, tk.END)

    # Show Output in Right Frame
    clear_right_frame()
    output_text = (
        f"Record Saved Successfully!\n\n"
        f"Name: {s_name}\n"
        f"Roll No: {r_no}\n"
        f"Class: {s_class}\n"
        f"Total Marks: {total} / 500\n"
        f"Percentage: {percentage:.2f}%\n"
        f"Result: {result}"
    )
    lbl = tk.Label(right_frame, text=output_text, font=("Arial", 12, "bold"), bg="lightyellow", justify="left")
    lbl.pack(pady=20, padx=20)


# Step 4: Get Result Function
def show_get_result():
    clear_right_frame()

    title = tk.Label(right_frame, text="Search Result", font=("Arial", 14, "bold"), bg="lightyellow")
    title.pack(pady=10)

    lbl_prompt = tk.Label(right_frame, text="Enter Roll No:", font=("Arial", 11), bg="lightyellow")
    lbl_prompt.pack()

    search_entry = tk.Entry(right_frame, font=("Arial", 11))
    search_entry.pack(pady=5)

    res_lbl = tk.Label(right_frame, text="", font=("Arial", 11), bg="lightyellow", justify="left")
    res_lbl.pack(pady=10)

    def search_action():
        s_roll = search_entry.get().strip()
        if not s_roll:
            messagebox.showerror("Error", "Enter Roll No to search!")
            return

        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb.active

        found = False
        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[0]) == str(s_roll):
                info = (
                    f"Name: {row[1]}\n"
                    f"Roll No: {row[0]}\n"
                    f"Class: {row[2]}\n"
                    f"Total Marks: {row[8]}\n"
                    f"Percentage: {row[9]}%\n"
                    f"Result: {row[10]}"
                )
                res_lbl.config(text=info, fg="black", font=("Arial", 11, "bold"))
                found = True
                break
        wb.close()

        if not found:
            res_lbl.config(text="std record not found", fg="red", font=("Arial", 12, "bold"))

    btn_search = tk.Button(right_frame, text="Search", bg="lightblue", font=("Arial", 10, "bold"), command=search_action)
    btn_search.pack(pady=5)


# Step 5: Show All Records Function (Using Treeview Table)
def show_all_records():
    clear_right_frame()

    title = tk.Label(right_frame, text="All Student Records", font=("Arial", 14, "bold"), bg="lightyellow")
    title.pack(pady=10)

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active

    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            data.append(row)
    wb.close()

    if not data:
        tk.Label(right_frame, text="No records found in Excel.", bg="lightyellow", font=("Arial", 11)).pack(pady=20)
        return

    # Sort data by Roll Number
    data.sort(key=lambda x: int(x[0]) if str(x[0]).isdigit() else str(x[0]))

    # Define Treeview Table Columns
    columns = ("Roll No", "Name", "Class", "Total", "%", "Result")
    tree = ttk.Treeview(right_frame, columns=columns, show="headings", height=14)

    # Define Column Headings and Widths
    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=80, anchor="center")

    # Insert Sorted Rows into Treeview Table
    for row in data:
        # row indexes from excel: 0=RollNo, 1=Name, 2=Class, 8=Total, 9=Percentage, 10=Result
        tree.insert("", tk.END, values=(row[0], row[1], row[2], row[8], f"{row[9]}%", row[10]))

    tree.pack(fill="both", expand=True, padx=10, pady=10)


# Step 6: GUI Main Window Setup
window = tk.Tk()
window.title("Student Management System")
window.geometry("900x550")
window.config(bg="lightpink")

# Header Label
header = tk.Label(window, text="Student Management System", font=("Georgia", 18, "bold"), bg="lightblue", height=2)
header.pack(fill="x", pady=5)

# Left Frame for Inputs
left_frame = tk.Frame(window, bg="lightpink")
left_frame.pack(side="left", padx=20, pady=10, anchor="n")

# Right Frame for Displaying Outputs
right_frame = tk.Frame(window, bg="lightyellow", bd=2, relief="sunken")
right_frame.pack(side="right", padx=15, pady=10, fill="both", expand=True)

# Labels & Text Entries (Left Frame)
tk.Label(left_frame, text="Name:", bg="lightpink", font=("Arial", 10, "bold")).grid(row=0, column=0, sticky="e", pady=3)
name_entry = tk.Entry(left_frame)
name_entry.grid(row=0, column=1, pady=3)

tk.Label(left_frame, text="Roll No:", bg="lightpink", font=("Arial", 10, "bold")).grid(row=1, column=0, sticky="e", pady=3)
roll_entry = tk.Entry(left_frame)
roll_entry.grid(row=1, column=1, pady=3)

tk.Label(left_frame, text="Class:", bg="lightpink", font=("Arial", 10, "bold")).grid(row=2, column=0, sticky="e", pady=3)
class_entry = tk.Entry(left_frame)
class_entry.grid(row=2, column=1, pady=3)

tk.Label(left_frame, text="Python Marks:", bg="lightpink", font=("Arial", 10, "bold")).grid(row=3, column=0, sticky="e", pady=3)
sub1_entry = tk.Entry(left_frame)
sub1_entry.grid(row=3, column=1, pady=3)

tk.Label(left_frame, text="Java Marks:", bg="lightpink", font=("Arial", 10, "bold")).grid(row=4, column=0, sticky="e", pady=3)
sub2_entry = tk.Entry(left_frame)
sub2_entry.grid(row=4, column=1, pady=3)

tk.Label(left_frame, text="Math Marks:", bg="lightpink", font=("Arial", 10, "bold")).grid(row=5, column=0, sticky="e", pady=3)
sub3_entry = tk.Entry(left_frame)
sub3_entry.grid(row=5, column=1, padx=5, pady=3)

tk.Label(left_frame, text="Dart Marks:", bg="lightpink", font=("Arial", 10, "bold")).grid(row=6, column=0, sticky="e", pady=3)
sub4_entry = tk.Entry(left_frame)
sub4_entry.grid(row=6, column=1, pady=3)

tk.Label(left_frame, text="SQL Marks:", bg="lightpink", font=("Arial", 10, "bold")).grid(row=7, column=0, sticky="e", pady=3)
sub5_entry = tk.Entry(left_frame)
sub5_entry.grid(row=7, column=1, pady=3)

# Buttons (Left Frame)
tk.Button(left_frame, text="Save", width=15, bg="lightgreen", font=("Arial", 10, "bold"), command=save_data).grid(row=8, column=0, columnspan=2, pady=10)
tk.Button(left_frame, text="Get Result", width=15, bg="khaki", font=("Arial", 10, "bold"), command=show_get_result).grid(row=9, column=0, columnspan=2, pady=5)
tk.Button(left_frame, text="Show All Records", width=15, bg="cyan", font=("Arial", 10, "bold"), command=show_all_records).grid(row=10, column=0, columnspan=2, pady=5)

window.mainloop()
